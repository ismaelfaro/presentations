"""
Notas del ponente, reescritas sobre el contenido y el orden actuales de Ismael_Faro_Presentacion.pptx (34 slides).

  uv run --with python-pptx python talks/26.Madrid.Global-Tech-forum/notas_ponente.py

Genera:
  - Ismael_Faro_Presentacion_notas.pptx  (copia del .pptx con las notas nuevas; el original no se toca)
  - notas_ponente.md                      (todas las notas en un documento, para repasar o imprimir)
  - site/26.Madrid.Global-Tech-forum/slides.json   (actualiza las notas que ve la vista ponente de la web)

Formato de cada nota: [minuto · BLOQUE], ideas clave en viñetas (cifras y frases ancla) y fuentes al final.
Tiempos pensados para 40 min: la charla acaba hacia el 38:30 y deja ~1,5 min de colchón.
"""

import json
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).parent
SRC = ROOT / "Ismael_Faro_Presentacion.pptx"
OUT = ROOT / "Ismael_Faro_Presentacion_notas.pptx"

NOTES = {
1: """[0:00 · APERTURA]
• Gracias, Ester: tú la máquina, yo de dónde viene.
• La IA es un paso más en una cadena muy larga. No el primero.""",

2: """[1:00 · SILICIO]
• Vaso y móvil: mismo material. Arena.
• Ordenado = cristal; desordenado = vidrio. Aislado por Berzelius, 1824.""",

3: """[2:00 · VIDRIO]
• +4.000 años dándole forma. Jarrón egipcio, 1500 a. C.
• Primero se miraba el vidrio; luego, a través de él.""",

4: """[3:00 · LA LENTE]
• 1608 Lippershey · 1610 Galileo, lunas de Júpiter · 1665 Hooke, «células».
• Hacia fuera, el cosmos. Hacia dentro, la vida.
• ANCLA 1: «Una herramienta no responde preguntas: cambia qué preguntas podemos hacer.» Pausa.""",

5: """[4:30 · CAJA NEGRA]
• Usamos lentes siglos antes de entender la luz.
• Hoy, otra caja: sabemos qué entra y qué sale, no todo lo de dentro.""",

6: """[5:30 · LA IA, NUEVA LENTE]
• Hacia dentro: AlphaFold, proteínas, Nobel de Química 2024.
• Hacia fuera: WeatherNext, tiempo y clima.
• Mira patrones, lenguaje, decisiones.
Fuentes: https://www.nobelprize.org/prizes/chemistry/2024/summary/ · https://deepmind.google/science/weathernext/""",

7: """[7:00 · LA BOMBILLA]
• 1879 Edison/Swan: frágil, cara, unas horas.
• 1881 Savoy · 4 sept. 1882 Pearl Street: un barrio de Manhattan. Menos de 3 años.
• Edison no inventó la bombilla: inventó la red. Guardad esta idea.
Fuentes: https://en.wikipedia.org/wiki/Savoy_Theatre · https://en.wikipedia.org/wiki/Pearl_Street_Station""",

8: """[9:00 · CAMBIO QUE YA VIVIMOS]
• Internet: ~30.000 M dispositivos, 3-4 por persona.
• Smartphone: juguete en 2007, hoy más de media humanidad.
• ChatGPT: 100 M usuarios en 2 meses.
• Misma locura que la bombilla. Un paso más.
Fuentes: https://www.cisco.com/c/en/us/solutions/collateral/executive-perspectives/annual-internet-report/white-paper-c11-741490.html · https://www.reuters.com/technology/chatgpt-sets-record-fastest-growing-user-base-analyst-note-2023-02-01/""",

9: """[11:00 · TITULARES]
• Hoy: «la IA acabará con todo». No es nuevo.
• 1865 ley de la bandera roja · 1889 John Feeks, Nueva York.
• Miedos reales; la respuesta fue dirigir, no prohibir.
Fuentes: https://en.wikipedia.org/wiki/Locomotive_Acts · https://en.wikipedia.org/wiki/War_of_the_currents""",

10: """[12:30 · OFICIOS]
• El sereno, el farolero: desaparecen. Llegan electricistas e ingenieros.
• Perdemos hábitos, ganamos otros.
Imagen: https://es.wikipedia.org/wiki/Sereno_%28oficio%29""",

11: """[13:30 · DATOS DE EMPLEO]
• 1 de cada 3 trabajadores en ocupaciones muy expuestas. Madrid, la que más.
• Desde 2022 ese empleo crece por encima de la media: tareas, no puestos.
• Pero ofertas junior −33 % en un año. El cambio empieza por quienes empiezan: formación.
Fuentes: https://www.bbvaresearch.com/publicaciones/espana-observatorio-trimestral-del-mercado-de-trabajo-2t2026/ · https://dplnews.com/espana-la-ia-transforma-el-acceso-al-empleo-tecnologico-las-ofertas-para-programadores-junior-caen-un-33-mientras-los-perfiles-senior-crecen-un-13/""",

12: """[15:30 · MIEDOS]
• Adopción desigual · empleo · concentración de poder · ética y sesgos.
• Mañana, módulo de gobernanza y ética.
• ANCLA 2: «La herramienta no decide. Decidimos nosotros.» Pausa.""",

13: """[16:30 · MODELO Y AGENTE]
• Modelo = bombilla. Agente (harness) = red: herramientas, datos, memoria, permisos.
• El valor estará en quien sepa construir la red: organización y personas.
• Opcional: «esta presentación la muevo con un agente que corre en el portátil».""",

14: """[18:30 · LO QUE VIENE]
• Más rápidos · más expertos · capaces de mejorarse.
• AlphaEvolve (2025): matrices 4×4 con 48 multiplicaciones; récord de 1969.
Fuentes: https://deepmind.google/discover/blog/alphaevolve-a-gemini-powered-coding-agent-for-designing-advanced-algorithms/ · https://arxiv.org/abs/2506.13131""",

15: """[21:00 · IA QUE INVESTIGA]
• Migraciones de código: de semanas a días (Qiskit).
• Circuitos: 1.917 → 1.692 CNOT, 150 → 90 capas.
• Nuestro paper: 465 códigos de corrección de errores, ~140 h, ~400 $.
Fuente: https://arxiv.org/abs/2606.02418""",

16: """[23:00 · LA IA COMO LLAVE]
• De patrones a conceptos.
• Tres puertas: robótica, cuántica, biología.""",

17: """[24:00 · ROBÓTICA]
• Juegos de Robots Humanoides de Pekín: corre raro, nos reímos.
• Es nuestra bombilla de 1879.
Fuente: https://en.wikipedia.org/wiki/World_Humanoid_Robot_Games""",

18: """[25:00 · WAYMO]
• Su Pearl Street: +500.000 viajes autónomos por semana.
• Ya no es demo; es servicio.
Fuente: https://9to5google.com/2026/04/29/alphabet-q1-2026-earnings/""",

19: """[26:00 · BRAZO SO-100, 2024]
• Robótica como juguete para aprender: impreso en 3D, precio de un móvil.
Fuente: https://github.com/TheRobotStudio/SO-ARM100""",

20: """[26:40 · LEROBOT]
• Hugging Face: modelos, datos y robots abiertos para cualquier estudiante.
Fuente: https://github.com/huggingface/lerobot""",

21: """[27:20 · MICRODUCK, 2026]
• Pato bípedo de 25 cm, ~400 $. Del laboratorio al pupitre.
Fuente: https://techcrunch.com/2026/08/27/hugging-face-is-selling-a-cute-399-open-source-duck-robot-microduck/""",

22: """[28:00 · CUÁNTICA]
• Otra forma de procesar información, no un ordenador más rápido.
• IA y cuántica se necesitan. RIKEN: procesador IBM + Fugaku entero.
• Enrique contará más.
Fuente: https://arxiv.org/abs/2511.00224""",

23: """[29:00 · BOSCH]
• Química y materiales con algoritmos híbridos (consorcio QUTAC).
Fuente: https://arxiv.org/abs/2301.11838""",

24: """[29:40 · MODERNA]
• Acuerdo con IBM (2023): IA generativa + cuántica para ARNm.
• Ya: plegamiento de ARNm en procesadores de hasta 156 cúbits.
Fuentes: https://newsroom.ibm.com/2023-04-20-Moderna-and-IBM-to-Explore-Quantum-Computing-and-Generative-AI-for-mRNA-Science · https://arxiv.org/abs/2505.05782""",

25: """[30:10 · E.ON]
• Optimizar el transporte de energía en la red eléctrica.
Fuente: https://arxiv.org/abs/2403.17495""",

26: """[30:40 · HSBC]
• Bonos reales: hasta +34 % en predecir qué órdenes se ejecutan.
• El ruido cuántico ayudó.
Fuentes: https://arxiv.org/abs/2509.17715 · https://www.hsbc.com/news-and-views/news/media-releases/2025/hsbc-demonstrates-worlds-first-known-quantum-enabled-algorithmic-trading-with-ibm""",

27: """[31:30 · NEURONAS Y PONG]
• 2022: neuronas sobre silicio aprenden Pong.
• 2025: CL1, primer ordenador biológico comercial (MWC Barcelona).
Fuentes: https://doi.org/10.1016/j.neuron.2022.09.001 · https://corticallabs.com/cl1""",

28: """[32:30 · HONGOS]
• Círculo o cruz: el micelio crece distinto. «Reconoce» la forma sin cerebro.
• La internet del bosque.
Fuente: https://doi.org/10.1016/j.funeco.2024.101387""",

29: """[33:30 · FRESNEL]
• Vuelvo a la luz. 1822: anillos de cristal; 1823, faro de Cordouan.
• Guiño: «con mi apellido, tenía que pasar por los faros».
Fuente: https://cordouan.culture.gouv.fr/en/cordouans-first-lens""",

30: """[34:00 · EL FARO]
• Una luz pequeña, vista a decenas de kilómetros.
• La lente no crea la luz: la dirige.""",

31: """[34:45 · DIRIGIR]
• Nada de esto se frenó; quien frenó, se quedó atrás.
• ANCLA 3: «La innovación no hay que frenarla. Hay que dirigirla.» Pausa larga.""",

32: """[35:30 · ESPAÑA]
• Actuador, no observador (hilo de Virkkunen).
• Talento · reglas que dirigen · inversión.
• La ventana es ahora.""",

33: """[36:45 · ANFAIA]
• Jóvenes que crean y dirigen IA, no solo la consumen.
• Becas de verano con agentes: Smart Health, ClimaSafe, MalariaSentinel, SkillNet.
• Los primeros electricistas llevaron la luz a cada casa.
Fuente: https://anfaia.org/""",

34: """[38:00 · FINAL]
• Moldeamos el cristal y la luz; ahora toca moldear la IA y dirigir los agentes.
• La herramienta la construimos nosotros. La decisión, también.
• Gracias. Paso a Enrique.""",
}


def main():
    prs = Presentation(SRC)
    slides = list(prs.slides)
    assert len(slides) == len(NOTES), f"el .pptx tiene {len(slides)} slides y hay {len(NOTES)} notas"
    for i, s in enumerate(slides, 1):
        s.notes_slide.notes_text_frame.text = NOTES[i]
    prs.save(OUT)
    print(f"{OUT.name}: {len(slides)} notas escritas")

    md = ["# Notas del ponente", "",
          "*Cómo la IA transformará nuestras vidas y trabajo en la próxima década* · Foro Tecnológico Global · 40 min", "",
          "Frases ancla: (1) «Una herramienta no responde preguntas: cambia qué preguntas podemos hacer.» "
          "(2) «La herramienta no decide. Decidimos nosotros.» (3) «La innovación no hay que frenarla. Hay que dirigirla.»", ""]
    for i, s in enumerate(slides, 1):
        md += [f"## {i}", "", NOTES[i].replace("\n", "\n\n"), ""]
    (ROOT / "notas_ponente.md").write_text("\n".join(md))
    print("notas_ponente.md escrito")

    web = ROOT.parent.parent / "site" / ROOT.name / "slides.json"
    data = json.loads(web.read_text())
    if len(data) == len(NOTES):
        for d in data:
            d["notes"] = NOTES[d["n"]]
        web.write_text(json.dumps(data, ensure_ascii=False, indent=1))
        print(f"{web.relative_to(ROOT.parent.parent)} actualizado")
    else:
        print(f"{web.name} tiene {len(data)} slides: no se actualiza")


if __name__ == "__main__":
    main()
